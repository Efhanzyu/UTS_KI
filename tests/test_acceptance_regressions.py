"""Regression coverage for the measured acceptance criteria."""
import io
import re
from pathlib import Path

import pytest
from config import Config, resolve_secret_key
from app.routes import api as api_routes


@pytest.mark.parametrize("plaintext", [
    " leading", "trailing ", "\nstarts with newline", "ends with newline\n",
    "line one\nline two", "tabs\tand symbols !@#$%", "  \n  ",
])
def test_text_api_roundtrip_preserves_exact_whitespace(client, monkeypatch, sample_password, plaintext):
    monkeypatch.setattr(Config, "PBKDF2_ITERATIONS", 1000)
    encrypted = client.post("/api/encrypt/text", json={
        "plaintext": plaintext, "password": sample_password, "algorithm": "AES-256-GCM"
    })
    assert encrypted.status_code == 200
    ciphertext = encrypted.get_json()["data"]["ciphertext_b64"]
    decrypted = client.post("/api/decrypt/text", json={
        "ciphertext_b64": ciphertext, "password": sample_password
    })
    assert decrypted.status_code == 200
    assert decrypted.get_json()["data"]["plaintext"] == plaintext


def test_empty_password_is_rejected_before_encryption_service(client, monkeypatch):
    def fail_if_called(*args, **kwargs):
        pytest.fail("encryption service must not run without a password")
    monkeypatch.setattr(api_routes.encryption_service, "encrypt_text", fail_if_called)
    response = client.post("/api/encrypt/text", json={"plaintext": "hello", "password": ""})
    assert response.status_code == 400
    assert "password" in response.get_json()["message"].lower()


def test_wrong_password_and_tampered_file_are_rejected_without_download(client, monkeypatch, sample_password):
    monkeypatch.setattr(Config, "PBKDF2_ITERATIONS", 1000)
    source = b"%PDF-1.4\nroundtrip check\n%%EOF"
    encrypted = client.post("/api/encrypt/file", data={
        "file": (io.BytesIO(source), "sample.pdf"),
        "password": sample_password,
        "algorithm": "AES-256-GCM",
    }, content_type="multipart/form-data")
    assert encrypted.status_code == 200
    ciphertext = encrypted.data

    wrong = client.post("/api/decrypt/file", data={
        "file": (io.BytesIO(ciphertext), "sample.pdf.enc"), "password": "incorrect-password"
    }, content_type="multipart/form-data")
    assert wrong.status_code == 400
    assert wrong.mimetype == "application/json"

    modified = bytearray(ciphertext)
    modified[-1] ^= 1
    tampered = client.post("/api/decrypt/file", data={
        "file": (io.BytesIO(modified), "tampered.enc"), "password": sample_password
    }, content_type="multipart/form-data")
    assert tampered.status_code == 400
    assert tampered.mimetype == "application/json"
    assert "gagal" in tampered.get_json()["message"].lower()


def test_production_secret_is_required_and_development_secret_is_ephemeral():
    with pytest.raises(RuntimeError, match="SECRET_KEY"):
        resolve_secret_key(None, "production")
    first = resolve_secret_key(None, "development")
    second = resolve_secret_key(None, "development")
    assert len(first) >= 64
    assert first != second


def test_env_example_does_not_contain_a_secret_and_env_is_ignored():
    example = Path(".env.example").read_text(encoding="utf-8")
    assert re.search(r"(?m)^SECRET_KEY\s*=\s*$", example)
    assert re.search(r"(?m)^SUPABASE_SERVICE_ROLE_KEY\s*=\s*$", example)
    assert "SUPABASE_STORAGE_BUCKET=brankas-files" in example
    ignore = Path(".gitignore").read_text(encoding="utf-8")
    assert ".env" in ignore and ".env.*" in ignore and "!.env.example" in ignore


def test_supabase_config_reads_environment_variable_names_without_embedded_keys():
    source = Path("config.py").read_text(encoding="utf-8")
    assert 'os.environ.get("SUPABASE_URL"' in source
    assert 'os.environ.get("SUPABASE_SERVICE_ROLE_KEY"' in source
    assert 'os.environ.get("SUPABASE_STORAGE_BUCKET"' in source
    assert "os.environ.get(\"sb_" not in source


def test_frontend_does_not_persist_password_or_put_it_in_url():
    scripts = chr(10).join(path.read_text(encoding="utf-8") for path in Path("static/js").glob("*.js"))
    assert "localStorage" not in scripts
    assert "sessionStorage" not in scripts
    marker = "URLSearchParams(window.location.search)"
    assert marker in scripts
    query_code = scripts[scripts.index(marker):scripts.index(marker) + 220]
    assert "password" not in query_code.lower()


def test_runtime_sources_have_no_literal_secret_or_benchmark_password():
    sources = "\n".join(
        path.read_text(encoding="utf-8")
        for folder in (Path("app"),)
        for path in folder.rglob("*.py")
    ) + "\n" + Path("config.py").read_text(encoding="utf-8")
    assert "_BENCHMARK_PASSWORD" not in sources
    assert not re.search(r"(?:SECRET_KEY|_BENCHMARK_PASSWORD)\s*=\s*['\"][^'\"]{8,}['\"]", sources)


def test_benchmark_reports_kdf_cipher_entropy_histogram_and_avalanche(client, monkeypatch):
    monkeypatch.setattr(Config, "PBKDF2_ITERATIONS", 1000)
    monkeypatch.setattr(Config, "BENCHMARK_KDF_ITERATIONS", 1000)
    response = client.post("/api/benchmark", json={"sizes_kb": [1], "runs": 1})
    assert response.status_code == 200
    data = response.get_json()["data"]
    assert {row["algorithm"] for row in data["time_benchmark"]} == {
        "AES-256-GCM", "ChaCha20-Poly1305"
    }
    assert data["kdf_benchmark"]["mean_ms"] >= 0
    assert len(data["time_benchmark"]) == 2
    for row in data["time_benchmark"]:
        assert row["size_bytes"] == 1024
        assert row["avg_encryption_ms"] >= 0
        assert row["avg_decryption_ms"] >= 0
    assert len(data["entropy_comparison"]) == 2
    assert len(data["histogram"]) == 2
    for row in data["histogram"]:
        assert len(row["plaintext_histogram"]) == 256
        assert len(row["ciphertext_histogram"]) == 256
    assert len(data["avalanche"]) == 2
    for row in data["avalanche"]:
        flip = row["plaintext_bit_flip"]
        assert row["avalanche_percentage"] == round(
            flip["changed_bits"] * 100 / flip["total_bits"], 2
        )


def test_benchmark_rejects_runs_outside_configured_range(client):
    response = client.post("/api/benchmark", json={"sizes_kb": [1], "runs": 101})
    assert response.status_code == 400


def test_invalid_base64_and_container_are_rejected_without_traceback(client):
    for ciphertext in ("not_base64***", "bm90LWEtYnJhbmtzLWNvbnRhaW5lcg=="):
        response = client.post("/api/decrypt/text", json={
            "ciphertext_b64": ciphertext, "password": "any-password"
        })
        assert response.status_code == 400
        assert response.get_json()["success"] is False
        assert "Traceback" not in response.get_json()["message"]


def test_unsupported_algorithm_is_rejected(client):
    response = client.post("/api/encrypt/text", json={
        "plaintext": "valid text", "password": "password", "algorithm": "AES-256-CBC"
    })
    assert response.status_code == 400
    assert response.get_json()["success"] is False


@pytest.mark.parametrize("path", ["/", "/encrypt", "/decrypt", "/text", "/benchmark"])
def test_primary_pages_return_html_success(client, path):
    response = client.get(path)
    assert response.status_code == 200
    assert response.mimetype == "text/html"


def test_avalanche_copy_and_exports_do_not_claim_a_fixed_ideal():
    sources = [
        Path("templates/benchmark.html").read_text(encoding="utf-8"),
        Path("static/js/benchmark.js").read_text(encoding="utf-8"),
        Path("scripts/run_benchmark.py").read_text(encoding="utf-8"),
        Path("scripts/export_results.py").read_text(encoding="utf-8"),
    ]
    joined = chr(10).join(sources)
    assert "Target Ideal" not in joined
    assert "50.0%" not in joined
    assert "calculate_avalanche_test_only" in Path("app/services/benchmark_service.py").read_text(encoding="utf-8")


def test_benchmark_csv_and_xlsx_export_measured_kdf_and_both_avalanche_experiments(tmp_path, monkeypatch):
    from app.services import benchmark_service
    from scripts.export_results import export_csv, export_xlsx

    monkeypatch.setattr(Config, "PBKDF2_ITERATIONS", 1000)
    monkeypatch.setattr(Config, "BENCHMARK_KDF_ITERATIONS", 1000)
    result = benchmark_service.run_full_benchmark(sizes_kb=[1], runs=1)
    csv_path = tmp_path / "measurement.csv"
    xlsx_path = tmp_path / "measurement.xlsx"
    export_csv(result, csv_path)
    export_xlsx(result, xlsx_path)

    csv_text = csv_path.read_text(encoding="utf-8")
    assert "KDF PBKDF2" in csv_text
    assert "Plaintext bit" in csv_text and "Key bit" in csv_text
    assert "Target Ideal" not in csv_text

    import openpyxl
    workbook = openpyxl.load_workbook(xlsx_path, read_only=True)
    assert "PBKDF2 KDF" in workbook.sheetnames
    avalanche = workbook["Avalanche Effect"]
    assert avalanche.max_row == 7
    assert "Plaintext bit" in {avalanche.cell(row=i, column=2).value for i in range(4, 8)}
    assert "Key bit" in {avalanche.cell(row=i, column=2).value for i in range(4, 8)}
    workbook.close()
