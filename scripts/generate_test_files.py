"""
Script untuk membuat 10 file uji sintetis untuk pengujian enkripsi dan benchmark.
Mencakup dokumen akademik, gambar, dataset, dan arsip.
"""
import os
import zipfile
import json
import csv
from io import BytesIO
from PIL import Image, ImageDraw
import docx
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "test_files")


def ensure_output_dir():
    os.makedirs(OUTPUT_DIR, exist_ok=True)


def create_txt():
    path = os.path.join(OUTPUT_DIR, "01_catatan_kuliah.txt")
    content = """Catatan Kuliah: Keamanan Informasi & Kriptografi
Topik: Algoritma Kriptografi Modern (AES-256-GCM & ChaCha20-Poly1305)

1. Confidentiality: Melindungi kerahasiaan data menggunakan cipher simetris 256-bit.
2. Integrity: Menjamin data tidak dimodifikasi menggunakan Message Authentication Code (MAC).
3. Authenticity: Memastikan sumber dan keabsahan ciphertext sebelum didekripsi.
4. Key Derivation: Menggunakan PBKDF2-HMAC-SHA256 dengan salt acak 32 bytes untuk mitigasi brute force.

Kerahasiaan tugas dan data akademik terjaga dengan aman.
"""
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"Created: {path}")


def create_csv():
    path = os.path.join(OUTPUT_DIR, "02_daftar_nilai_tugas.csv")
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["NIM", "Nama Mahasiswa", "Tugas 1", "Tugas 2", "UTS", "Status"])
        for i in range(1, 51):
            writer.writerow([f"20260801{i:03d}", f"Mahasiswa {i}", 85 + (i % 15), 80 + (i % 20), 88 + (i % 12), "LULUS"])
    print(f"Created: {path}")


def create_json():
    path = os.path.join(OUTPUT_DIR, "03_konfigurasi_akademik.json")
    data = {
        "mata_kuliah": "Keamanan Informasi",
        "tahun_ajaran": "2026/2027",
        "modul": [
            {"id": 1, "nama": "Kriptografi Simetris", "algoritma": ["AES-256-GCM", "ChaCha20-Poly1305"]},
            {"id": 2, "nama": "Key Derivation Function", "algoritma": ["PBKDF2-HMAC-SHA256"]},
            {"id": 3, "nama": "Integritas Data", "metode": ["AEAD Authentication Tag"]}
        ],
        "pengaturan_keamanan": {
            "min_password_len": 1,
            "max_password_len": 256,
            "salt_size_bytes": 32,
            "nonce_size_bytes": 12
        }
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=4)
    print(f"Created: {path}")


def create_pdf():
    path = os.path.join(OUTPUT_DIR, "04_laporan_praktikum.pdf")
    c = canvas.Canvas(path, pagesize=letter)
    c.setFont("Helvetica-Bold", 16)
    c.drawString(100, 750, "LAPORAN PRAKTIKUM KEAMANAN INFORMASI")
    c.setFont("Helvetica", 12)
    c.drawString(100, 720, "Mata Kuliah: Keamanan Informasi")
    c.drawString(100, 700, "Topik: Implementasi Brankas File Tugas Kuliah")
    c.drawString(100, 680, "Algoritma: AES-256-GCM & ChaCha20-Poly1305")
    c.drawString(100, 650, "Dokumen ini merupakan file PDF pengujian enkripsi dan dekripsi.")
    c.drawString(100, 630, "Isi dokumen terproteksi secara matematis melalui enkripsi terotentikasi.")
    c.save()
    print(f"Created: {path}")


def create_docx():
    path = os.path.join(OUTPUT_DIR, "05_makalah_keamanan.docx")
    doc = docx.Document()
    doc.add_heading("Makalah: Evaluasi Performa Kriptografi Modern", 0)
    p = doc.add_paragraph("Analisis komparatif antara AES-256-GCM dan ChaCha20-Poly1305 ")
    p.add_run("dalam pengamanan file tugas dan repositori akademik mahasiswa.").bold = True
    doc.add_heading("1. Pendahuluan", level=1)
    doc.add_paragraph(
        "Keamanan dokumen mahasiswa menjadi prioritas di era digital. "
        "Penggunaan enkripsi authenticated menjamin confidentiality sekaligus integrity."
    )
    doc.save(path)
    print(f"Created: {path}")


def create_png():
    path = os.path.join(OUTPUT_DIR, "06_diagram_arsitektur.png")
    img = Image.new("RGB", (600, 400), color=(15, 29, 53))
    draw = ImageDraw.Draw(img)
    draw.rectangle([50, 50, 550, 350], outline=(59, 130, 246), width=3)
    draw.text((120, 180), "Brankas File Tugas Kuliah - Diagram Arsitektur", fill=(226, 232, 240))
    img.save(path)
    print(f"Created: {path}")


def create_jpg():
    path = os.path.join(OUTPUT_DIR, "07_foto_kegiatan.jpg")
    img = Image.new("RGB", (800, 600), color=(20, 35, 65))
    draw = ImageDraw.Draw(img)
    draw.rectangle([30, 30, 770, 570], outline=(16, 185, 129), width=4)
    draw.text((250, 290), "Dokumentasi Praktikum KI 2026", fill=(255, 255, 255))
    img.save(path, "JPEG", quality=90)
    print(f"Created: {path}")


def create_zip():
    path = os.path.join(OUTPUT_DIR, "08_source_code_arsip.zip")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("main.py", "print('Sistem Pengaman Berkas Terdistribusi')\n")
        zf.writestr("README.txt", "Arsip source code pendukung penelitian.\n")
    print(f"Created: {path}")


def create_1mb_bin():
    path = os.path.join(OUTPUT_DIR, "09_dataset_penelitian_1mb.bin")
    with open(path, "wb") as f:
        f.write(os.urandom(1024 * 1024))
    print(f"Created: {path}")


def create_10mb_bin():
    path = os.path.join(OUTPUT_DIR, "10_dokumen_skripsi_10mb.bin")
    with open(path, "wb") as f:
        f.write(os.urandom(10 * 1024 * 1024))
    print(f"Created: {path}")


def main():
    ensure_output_dir()
    print(f"Generating 10 test files into {OUTPUT_DIR}...")
    create_txt()
    create_csv()
    create_json()
    create_pdf()
    create_docx()
    create_png()
    create_jpg()
    create_zip()
    create_1mb_bin()
    create_10mb_bin()
    print("All 10 test files generated successfully.")


if __name__ == "__main__":
    main()
